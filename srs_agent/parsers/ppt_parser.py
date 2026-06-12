"""
PPT文档解析器
"""

from typing import Dict, Any
from pptx import Presentation


def parse_ppt(file_path: str) -> Dict[str, Any]:
    """
    解析PPT文件
    
    Args:
        file_path: PPT文件路径
        
    Returns:
        包含解析结果的字典
    """
    try:
        prs = Presentation(file_path)
        
        # 提取幻灯片内容
        slides_content = []
        for idx, slide in enumerate(prs.slides, 1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text)
            
            slides_content.append(f"Slide {idx}:\n" + '\n'.join(slide_text))
        
        content = '\n\n'.join(slides_content)
        
        # 提取元数据
        metadata = {
            'slide_count': len(prs.slides),
        }
        
        return {
            'content': content,
            'metadata': metadata,
            'success': True
        }
    except Exception as e:
        return {
            'content': '',
            'metadata': {},
            'success': False,
            'error': str(e)
        }
