"""
DOCX文档解析器
"""

from typing import Dict, Any
import docx


def parse_docx(file_path: str) -> Dict[str, Any]:
    """
    解析DOCX文件
    
    Args:
        file_path: DOCX文件路径
        
    Returns:
        包含解析结果的字典
    """
    try:
        doc = docx.Document(file_path)
        
        # 提取文本内容
        full_text = []
        for paragraph in doc.paragraphs:
            full_text.append(paragraph.text)
        
        content = '\n'.join(full_text)
        
        # 提取元数据
        metadata = {
            'paragraph_count': len(doc.paragraphs),
        }
        
        return {
            'content': content,
            'metadata': metadata,
            'success': True
        }
    except Exception as e:
        return {
            'content': '',
            'metadata': {},
            'success': False,
            'error': str(e)
        }
