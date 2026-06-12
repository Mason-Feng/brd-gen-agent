import os
from pathlib import Path
from typing import List, Dict, Optional
from brd_gen_agent.state.SRSState import SRSState
from brd_gen_agent.utils.log import logged_node, setup_logger

logger = setup_logger(__name__)

def parse_all_files(root_dir: str) -> List[Dict[str, str]]:
    """递归提取所有可识别的文本文件"""
    logger.info(f"开始解析目录: {root_dir}")
    documents = []
    file_count = 0
    for dirpath, _, filenames in os.walk(root_dir):
        for fname in filenames:
            # 跳过 Word 的临时文件（例如以 '~$' 开头的临时文件）
            if fname.startswith('~$'):
                logger.debug(f"跳过临时文件: {fname}")
                continue
            file_path = os.path.join(dirpath, fname)
            rel_path = os.path.relpath(file_path, root_dir)
            logger.debug(f"处理文件: {rel_path}")
            text = extract_text(file_path)
            if text is not None:
                documents.append({"path": rel_path, "content": text})
                file_count += 1
                logger.info(f"✓ 成功提取: {rel_path} ({len(text)} 字符)")
            else:
                logger.debug(f"✗ 跳过文件: {rel_path} (不支持的格式)")
    logger.info(f"目录解析完成: {root_dir} - 共提取 {file_count} 个文件")
    return documents

def extract_text(file_path: str) -> Optional[str]:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in ['.md', '.txt', '.py', '.java', '.js', '.ts', '.html', '.css', '.c', '.cpp', '.h', '.yaml', '.yml', '.json']:
            logger.debug(f"读取文本文件: {ext}")
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        elif ext == '.pdf':
            logger.debug(f"使用 PyPDF2 提取 PDF 内容")
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            logger.debug(f"PDF 提取完成: {len(reader.pages)} 页")
            return text
        elif ext == '.docx':
            logger.debug(f"使用 python-docx 提取 Word 内容")
            from docx import Document
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
            logger.debug(f"Word 提取完成: {len(doc.paragraphs)} 段落")
            return text
        elif ext == '.xlsx':
            logger.debug(f"使用 openpyxl 提取 Excel xlsx 内容")
            from openpyxl import load_workbook
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join("" if cell is None else str(cell) for cell in row)
                    if row_text.strip():
                        rows.append(row_text)
            text = "\n".join(rows)
            logger.debug(f"Excel xlsx 提取完成: {len(workbook.worksheets)} 个工作表")
            return text
        elif ext == '.xls':
            logger.debug(f"使用 xlrd 提取 Excel xls 内容")
            from xlrd import open_workbook
            workbook = open_workbook(file_path)
            rows = []
            for sheet in workbook.sheets():
                rows.append(f"Sheet: {sheet.name}")
                for row_idx in range(sheet.nrows):
                    row = [sheet.cell_value(row_idx, col_idx) for col_idx in range(sheet.ncols)]
                    row_text = "\t".join("" if cell is None else str(cell) for cell in row)
                    if row_text.strip():
                        rows.append(row_text)
            text = "\n".join(rows)
            logger.debug(f"Excel xls 提取完成: {len(workbook.sheets())} 个工作表")
            return text
        else:
            logger.debug(f"不支持的文件类型: {ext}")
            return None
    except Exception as e:
        logger.warning(f"文件解析失败: {file_path} - 错误: {str(e)}")
        return None


def save_parsed_documents_to_test_folder(documents: List[Dict[str, str]], output_dir: Optional[Path] = None) -> str:
    """Save parsed documents to a markdown file.

    By default saves to the current working directory. If `output_dir` is provided,
    it will be used instead.
    """
    if output_dir is None:
        output_dir = Path.cwd()
    else:
        output_dir = Path(output_dir)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / 'parsed_business_documents.md'
    with output_file.open('w', encoding='utf-8') as f:
        f.write('# 解析后的业务文档内容\n\n')
        for doc in documents:
            rel_path = doc.get('path', 'unknown')
            content = doc.get('content', '')
            ext = Path(rel_path).suffix.lower()

            f.write(f"## {rel_path}\n\n")
            f.write(f"- 文件类型: {ext or '未知'}\n\n")

            # 如果原文件是 Markdown，直接写入内容；否则放入代码块保留格式
            if ext in ['.md', '.markdown']:
                f.write(content.rstrip() + "\n\n")
            else:
                # 避免内容中出现 ```，选择合理的围栏长度
                fence = '```'
                if '```' in content:
                    fence = '````'
                f.write(f"{fence}\n")
                f.write(content.rstrip() + "\n")
                f.write(f"{fence}\n\n")
    logger.info(f"已将解析文档保存到: {output_file}")
    return str(output_file)


@logged_node
def parse_files_node(state: SRSState, config=None) -> SRSState:
    logger.info("=" * 50)
    logger.info("开始文件解析节点")
    
    logger.info(f"解析业务文档目录: {state['folder_path']}")
    docs = parse_all_files(state["folder_path"])
    logger.info(f"业务文档解析完成，共 {len(docs)} 个文件")
    
    if state.get("codebase_path"):
        logger.info(f"解析代码库目录: {state['codebase_path']}")
        codebase_docs = parse_all_files(state["codebase_path"])
        docs.extend(codebase_docs)
        logger.info(f"代码库解析完成，共 {len(codebase_docs)} 个文件，总计 {len(docs)} 个文件")
    else:
        logger.info("跳过代码库解析 (未提供路径)")
    
    # Save parsed documents to current working directory
    saved_path = save_parsed_documents_to_test_folder(docs, output_dir=Path.cwd())
    logger.info(f"已保存解析文档到当前目录: {saved_path}")

    logger.info(f"读取模板文件: {state['template_path']}")
    with open(state["template_path"], 'r', encoding='utf-8') as f:
        state["template_content"] = f.read()
    logger.info(f"模板读取完成 ({len(state['template_content'])} 字符)")
    
    logger.info("文件解析节点完成")
    logger.info("=" * 50)
    return {"documents": docs, "template_content": state["template_content"]}