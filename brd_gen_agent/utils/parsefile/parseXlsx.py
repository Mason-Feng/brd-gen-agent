import sys
import json
from pathlib import Path
from openpyxl import load_workbook

def parse_xlsx_to_dict(file_path):
    """
    将 xlsx 解析为结构化字典，便于序列化为 JSON 或 Markdown。
    返回格式：
    {
        "file": "文件名",
        "sheets": [
            {
                "name": "Sheet1",
                "headers": ["列1", "列2", ...],   # 第一行作为表头
                "rows": [
                    ["值1", "值2", ...],
                    ...
                ]
            },
            ...
        ]
    }
    """
    file_path = str(file_path)
    suffix = Path(file_path).suffix.lower()

    # 支持 xlsx（openpyxl）和 xls（xlrd）
    wb = None
    if suffix == '.xlsx' or suffix == '.xlsm' or suffix == '.xltx' or suffix == '.xltm':
        try:
            wb = load_workbook(file_path, data_only=True)
        except FileNotFoundError:
            raise FileNotFoundError(f"文件 '{file_path}' 未找到")
        except Exception as e:
            raise RuntimeError(f"打开 xlsx 文件出错: {e}")
    elif suffix == '.xls':
        try:
            # 延迟导入 xlrd，避免未安装时在 import 阶段失败
            import xlrd
            wb_xls = xlrd.open_workbook(file_path)
        except FileNotFoundError:
            raise FileNotFoundError(f"文件 '{file_path}' 未找到")
        except Exception as e:
            raise RuntimeError(f"打开 xls 文件出错: {e}")
        # 将 xlrd 风格的 workbook 对象标记为特殊处理
        wb = wb_xls
    else:
        raise RuntimeError(f"不支持的文件类型: {suffix}")

    result = {
        "file": file_path.split("/")[-1].split("\\")[-1],  # 仅文件名
        "sheets": []
    }

    # 根据 workbook 的类型分别处理 openpyxl 和 xlrd
    if suffix == '.xls':
        # xlrd workbook
        for sheet in wb.sheets():
            sheet_name = sheet.name
            nrows = sheet.nrows
            if nrows == 0:
                sheet_info = {"name": sheet_name, "headers": [], "rows": []}
                result["sheets"].append(sheet_info)
                continue

            # 第一行作为表头
            headers = [str(cell) if cell is not None else "" for cell in sheet.row_values(0)]
            rows = []
            for ridx in range(1, nrows):
                row = sheet.row_values(ridx)
                if all((cell is None or str(cell).strip() == '') for cell in row):
                    continue
                row_data = [str(cell) if cell is not None else "" for cell in row]
                rows.append(row_data)

            sheet_info = {"name": sheet_name, "headers": headers, "rows": rows}
            result["sheets"].append(sheet_info)
    else:
        # openpyxl workbook
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # 获取所有行数据（生成器转换为列表）
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                # 空工作表
                sheet_info = {"name": sheet_name, "headers": [], "rows": []}
                result["sheets"].append(sheet_info)
                continue

            # 第一行作为表头，其余为数据行
            headers = [str(cell) if cell is not None else "" for cell in all_rows[0]]
            rows = []
            for row in all_rows[1:]:
                # 避免整行全为空
                if all(cell is None for cell in row):
                    continue
                row_data = [str(cell) if cell is not None else "" for cell in row]
                rows.append(row_data)

            sheet_info = {
                "name": sheet_name,
                "headers": headers,
                "rows": rows
            }
            result["sheets"].append(sheet_info)

        wb.close()
    return result


def dict_to_markdown_tables(data_dict):
    """
    将上述字典结构转换为 Markdown 表格字符串（每个工作表一个表格）。
    """
    md_parts = []
    for sheet in data_dict["sheets"]:
        md_parts.append(f"### 工作表: {sheet['name']}\n")
        headers = sheet["headers"]
        rows = sheet["rows"]

        if not headers:
            md_parts.append("*（空工作表）*\n\n")
            continue

        # 表头行
        md_parts.append("| " + " | ".join(headers) + " |")
        # 分隔行
        md_parts.append("| " + " | ".join(["---"] * len(headers)) + " |")
        # 数据行
        for row in rows:
            # 填充不足列数（防止某些行缺少列）
            filled_row = row + [""] * (len(headers) - len(row))
            md_parts.append("| " + " | ".join(filled_row[:len(headers)]) + " |")
        md_parts.append("\n")

    return "\n".join(md_parts)


